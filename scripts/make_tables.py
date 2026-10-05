#!/usr/bin/env python3
"""Regenerate release inventory, boundary table and supplementary workbook.

S1 and S2 copy the public numeric records exactly. S3 combines recovered
author-reference membership aggregates for DrugBank/trueDDI and a reported
KEGG aggregate. Raw sources are not redistributed; the separate reference
rebuild script reproduces the first two comparisons from authorised copies.
Requires openpyxl for the workbook; the CSV files use only the standard library.
"""
import csv
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "tables"


def read(path):
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def write(name, rows):
    with (TABLES / name).open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    descriptions = {
        "ddi_enzyme_database.csv": "Primary unique unordered drug pairs; identifiers and source annotations; no clinical grades or withdrawn validation flags",
        "enzyme_pair_attribution.csv": "Long-format pair x source protein label x direction annotations",
        "cid_severity_bridge.csv": "Candidate PubChem CID to DDInter identifier bridge; 240 non-empty display names",
        "cid_drugbank_crosswalk.csv": "PubChem CID to DrugBank identifier bridge; identifier-only export",
        "enzyme_severity_stats.csv": "Nine tested symbols among 549 graded inhibition-attributed pairs; exploratory reuse",
        "enzyme_severity_stats_single_enzyme.csv": "Sensitivity analysis among 424 graded pairs with one inhibition-attributed symbol",
        "protein_annotation_audit.csv": "35 verified current human gene identities by UniProt accession; 33 source symbols agree and 2 conflict",
        "reference_membership_aggregate.csv": "Recovered author-reference membership counts: trueDDI and DrugBank; descriptive comparison",
        "severity_summary.csv": "Aggregate counts only: 93 Major, 419 Moderate, 50 Minor, 1338 Unknown, total 1900",
        "drugbank_validation_summary.csv": "Historical aggregate comparison; recovered 16316-pair author source reproduces 1172 testable / 73 matched pairs",
        "kegg_validation_summary.csv": "Historical reported aggregate comparison; original source snapshot unavailable",
        "liddi_coverage_comparison.csv": "Historical reported coverage comparison; original match intermediate unavailable",
    }
    inventory = [dict(file=f"data/{name}", rows=len(read(ROOT / "data" / name)), record=description)
                 for name, description in descriptions.items()]
    write("Table1_release_inventory.csv", inventory)
    sources = read(ROOT / "sources" / "sources.csv")
    boundaries = [dict(source=s["name"], role=s["role"], snapshot_status=s["snapshot_status"],
                       source_version=s["source_version"], original_retrieval_date=s["retrieval_date"],
                       reproducibility_boundary=s["limitation"])
                  for s in sources if s["kind"] == "upstream_resource"]
    write("Table2_source_reproducibility_boundaries.csv", boundaries)
    for source, target in (
        ("enzyme_severity_stats.csv", "SupplementaryTableS1_enzyme_statistics.csv"),
        ("enzyme_severity_stats_single_enzyme.csv", "SupplementaryTableS2_single_enzyme_sensitivity.csv"),
    ):
        write(target, read(ROOT / "data" / source))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    names = {"Table1_release_inventory.csv": "Table1_Inventory",
             "Table2_source_reproducibility_boundaries.csv": "Table2_Boundaries",
             "SupplementaryTableS1_enzyme_statistics.csv": "S1_Enzyme_Statistics",
             "SupplementaryTableS2_single_enzyme_sensitivity.csv": "S2_Single_Enzyme",
             "SupplementaryTableS3_external_aggregate_counts.csv": "S3_Reported_Aggregates"}
    for filename, title in names.items():
        data = read(TABLES / filename)
        ws = wb.create_sheet(title)
        ws.append(list(data[0]))
        for row in data:
            ws.append(list(row.values()))
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cells in ws.columns:
            letter = cells[0].column_letter
            ws.column_dimensions[letter].width = min(70, max(15, max(len(str(c.value or "")) for c in cells) + 2))
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(TABLES / "SupplementaryTables.xlsx")
    print("Generated Tables 1-2, supplementary S1/S2, and SupplementaryTables.xlsx; S3 retained as reported aggregate record.")


if __name__ == "__main__":
    main()
